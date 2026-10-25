import io
from datetime import datetime
from typing import Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database import db


def create_excel_report(target_date: Optional[str] = None) -> io.BytesIO:
    """
    Generates an Excel workbook (.xlsx) containing:
      1. Daily Attendance Summary (First IN, Last OUT, Duration, Status)
      2. Detailed Punch Logs (every IN/OUT event with timestamp)
    Returns an in-memory BytesIO buffer ready for download.
    """
    is_all_time = not target_date or target_date == "ALL"
    display_date = "All Recorded Dates" if is_all_time else target_date
    summary_data = db.get_daily_summary(target_date) if not is_all_time else []
    
    # If all-time is requested, we can gather summaries for all available dates
    if is_all_time:
        for d in db.get_available_dates():
            summary_data.extend(db.get_daily_summary(d))

    raw_logs = db.get_logs(date=target_date if not is_all_time else None, limit=2000)

    # Create Workbook
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Daily Attendance"
    ws_summary.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="0F172A")
    font_badge_in = Font(name="Segoe UI", size=10, bold=True, color="047857")
    font_badge_out = Font(name="Segoe UI", size=10, bold=True, color="B45309")
    
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_sub_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_row_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_row_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_badge_in = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_badge_out = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(border_style="medium", color="0F172A")
    header_border = Border(bottom=thick_bottom_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ================= SHEET 1: Daily Attendance Summary =================
    # Title Banner
    ws_summary.merge_cells("A1:K1")
    ws_summary["A1"] = "RFID ATTENDANCE REPORT"
    ws_summary["A1"].font = font_title
    ws_summary["A1"].alignment = align_left
    ws_summary.row_dimensions[1].height = 28

    ws_summary.merge_cells("A2:K2")
    ws_summary["A2"] = f"Report Date: {display_date}  |  Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} (IST)"
    ws_summary["A2"].font = font_subtitle
    ws_summary["A2"].alignment = align_left
    ws_summary.row_dimensions[2].height = 20

    # Summary KPI Stats Box (Row 4)
    stats = db.get_stats(target_date)
    kpis = [
        ("Total Active Today", len(summary_data)),
        ("Currently IN", stats["currently_in"]),
        ("Punched OUT", stats["checked_out"]),
        ("Total Registered", stats["total_registered_users"]),
    ]
    
    col_idx = 1
    for label, val in kpis:
        cell_lbl = ws_summary.cell(row=4, column=col_idx, value=label)
        cell_lbl.font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
        cell_lbl.fill = fill_row_alt
        cell_lbl.border = cell_border
        cell_lbl.alignment = align_center

        cell_val = ws_summary.cell(row=5, column=col_idx, value=val)
        cell_val.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
        cell_val.fill = fill_row_alt
        cell_val.border = cell_border
        cell_val.alignment = align_center
        col_idx += 1

    ws_summary.row_dimensions[4].height = 18
    ws_summary.row_dimensions[5].height = 24

    # Table Header (Row 7)
    headers = [
        "User ID",
        "Full Name",
        "Role",
        "Department",
        "RFID Card UID",
        "Date",
        "First Punch IN",
        "Last Punch OUT",
        "Work Duration",
        "Current Status",
        "Total Punches",
    ]

    for c_idx, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=7, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = header_border

    ws_summary.row_dimensions[7].height = 26

    # Populate Data
    start_row = 8
    if not summary_data:
        ws_summary.merge_cells("A8:K8")
        empty_cell = ws_summary["A8"]
        empty_cell.value = "No attendance records found for this date."
        empty_cell.font = Font(name="Segoe UI", size=11, italic=True, color="64748B")
        empty_cell.alignment = align_center
        ws_summary.row_dimensions[8].height = 30
    else:
        for idx, row in enumerate(summary_data):
            curr_row = start_row + idx
            ws_summary.row_dimensions[curr_row].height = 22
            row_fill = fill_row_alt if idx % 2 == 1 else fill_row_white

            status_str = row["status_label"]
            status_badge = row["status_badge"]

            vals = [
                (row["user_id"], align_center),
                (row["name"], align_left),
                (row["role"], align_left),
                (row["department"], align_left),
                (row["uid"], align_center),
                (row["date"], align_center),
                (row["first_in"] or "--:--:--", align_center),
                (row["last_out"] or "--:--:--", align_center),
                (row["duration_formatted"], align_center),
                (status_str, align_center),
                (row["total_punches"], align_center),
            ]

            for c_idx, (v, alignment) in enumerate(vals, 1):
                c = ws_summary.cell(row=curr_row, column=c_idx, value=v)
                c.font = font_data
                c.fill = row_fill
                c.border = cell_border
                c.alignment = alignment

                # Highlight Status column
                if c_idx == 10:
                    if status_badge == "IN":
                        c.fill = fill_badge_in
                        c.font = font_badge_in
                    elif status_badge == "OUT":
                        c.fill = fill_badge_out
                        c.font = font_badge_out

    # Auto-fit columns
    for col in ws_summary.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignore merged title row 1 & 2
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ================= SHEET 2: Detailed Punch Logs =================
    ws_logs = wb.create_sheet(title="Punch Logs")
    ws_logs.views.sheetView[0].showGridLines = True

    # Title
    ws_logs.merge_cells("A1:H1")
    ws_logs["A1"] = "RFID DETAILED PUNCH LOGS"
    ws_logs["A1"].font = font_title
    ws_logs["A1"].alignment = align_left
    ws_logs.row_dimensions[1].height = 28

    ws_logs.merge_cells("A2:H2")
    ws_logs["A2"] = f"Filter: {display_date}  |  Total Punch Events: {len(raw_logs)}"
    ws_logs["A2"].font = font_subtitle
    ws_logs["A2"].alignment = align_left
    ws_logs.row_dimensions[2].height = 20

    # Headers
    log_headers = [
        "Log ID",
        "User ID",
        "Full Name",
        "Card UID",
        "Action (Punch)",
        "Date",
        "Time",
        "Device ID",
    ]

    for c_idx, h in enumerate(log_headers, 1):
        cell = ws_logs.cell(row=4, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_sub_header
        cell.alignment = align_center
        cell.border = header_border

    ws_logs.row_dimensions[4].height = 24

    for idx, l in enumerate(raw_logs):
        curr_row = 5 + idx
        ws_logs.row_dimensions[curr_row].height = 20
        row_fill = fill_row_alt if idx % 2 == 1 else fill_row_white

        status = l.get("status", "IN")
        vals = [
            (f"#{l.get('id', '')}", align_center),
            (l.get("user_id", "UNREGISTERED"), align_center),
            (l.get("name", "Unregistered Card"), align_left),
            (l.get("uid", ""), align_center),
            (status, align_center),
            (l.get("date", ""), align_center),
            (l.get("time", ""), align_center),
            (l.get("device_id", "esp32_c3_gate1"), align_center),
        ]

        for c_idx, (v, alignment) in enumerate(vals, 1):
            c = ws_logs.cell(row=curr_row, column=c_idx, value=v)
            c.font = font_data
            c.fill = row_fill
            c.border = cell_border
            c.alignment = alignment

            if c_idx == 5:
                if status == "IN":
                    c.fill = fill_badge_in
                    c.font = font_badge_in
                else:
                    c.fill = fill_badge_out
                    c.font = font_badge_out

    # Auto-fit columns for logs sheet
    for col in ws_logs.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_logs.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
